"""
Junta as amostras citadas no resumo executivo e os filtros CT&I em uma planilha Excel única,
com uma aba por dataset e uma aba LEIA-ME que explica cada aba.

Uso:  python gerar_planilha_amostras.py [saida.xlsx]
      (padrão: TERRITORIOS_SECTI_Amostras_CTI.xlsx na raiz do projeto; requer openpyxl)
"""
import csv, re, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SAIDA = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "TERRITORIOS_SECTI_Amostras_CTI.xlsx"

# aba, fonte, arquivo, o que retorna, recorte, filtro CT&I, variáveis CT&I
ABAS = [
    ("IBGE_Populacao", "IBGE", "amostras/ibge_2.1_populacao_6579_BA.csv",
     "População residente estimada por município (SIDRA 6579)", "417 municípios × 2024–2026", "⬜ Sem variável", "—"),
    ("IBGE_PIB_VAB", "IBGE", "amostras/ibge_2.2_pib_vab_5938_BA.csv",
     "PIB e valor adicionado por setor, em mil R$ (SIDRA 5938), formato longo", "417 municípios × 2022–2023", "⬜ Sem variável", "—"),
    ("IBGE_PIB_per_capita_2023", "IBGE", "amostras/ibge_2.3_pib_per_capita_oficial_BA_2023.csv",
     "PIB e PIB per capita oficiais (base do PIB dos Municípios)", "417 municípios, 2023", "⬜ Sem variável", "—"),
    ("IBGE_Derivados", "IBGE", "amostras/ibge_2.3_2.5_2.6_derivados_amostra_BA.csv",
     "PIB per capita, densidade demográfica e PIB por km²", "15 municípios de referência, 2023", "⬜ Sem variável", "—"),
    ("IBGE_Area", "IBGE", "amostras/ibge_2.4_area_centroide_BA.csv",
     "Área (km²) e centroide de cada município", "417 municípios", "⬜ Sem variável", "—"),
    ("IBGE_Malha", "IBGE", "amostras/ibge_2.7_malha_amostra_BA.csv",
     "Metadados do polígono de cada município (o GeoJSON completo fica fora da planilha)", "30 municípios", "⬜ Sem variável", "—"),
    ("INEP_IES", "INEP", "amostras/inep_3.1_3.7_ies_BA_amostra.csv",
     "Instituições de ensino superior: tipo, rede, sede, docentes por titulação", "40 IES com mais doutores", "🔶 Não filtrado",
     "ies_tecnologica_if_cefet (IF/CEFET); QT_DOC_EX_DOUT"),
    ("INEP_Cursos", "INEP", "amostras/inep_3.2_a_3.6_cursos_BA_amostra.csv",
     "Cursos de graduação: área CINE, grau, modalidade, vagas, matrículas, ingressantes e concluintes",
     "Estratificada: 15 STEM pres. + 5 ampliado + 5 fora + 10 STEM EAD", "🏷️ Classificado",
     "nivel_cti ← CO_CINE_AREA_GERAL (05/06/07 núcleo; 08/09 ampliado)"),
    ("INEP_Cursos_Municipio", "INEP", "amostras/inep_cursos_presenciais_por_municipio_nivel_cti_amostra.csv",
     "Soma de cursos, vagas e matrículas presenciais por município e nível CT&I", "13 municípios de referência, só presencial",
     "🏷️ Classificado", "nivel_cti"),
    ("ANEEL_SIGA", "ANEEL", "amostras/aneel_4.1_4.2_siga_geracao_BA.csv",
     "Usinas de geração: tipo, fonte, fase, potência, coordenadas, município (texto)", "Todas as 1.163 usinas da BA",
     "🔶 Não filtrado", "SigTipoGeracao ∈ {EOL, UFV} (auxiliar)"),
    ("ANEEL_INDGER", "ANEEL", "amostras/aneel_4.3_indger_comercial_BA_amostra.csv",
     "Unidades consumidoras ativas e indicadores comerciais por município e mês", "COELBA, 1.500 linhas (2023–2026)",
     "⬜ Sem variável", "—"),
    ("ANEEL_SAMP", "ANEEL", "amostras/aneel_4.3_samp_mercado_COELBA_amostra.csv",
     "Mercado de energia por classe de consumo (kWh, kW, R$); sem município", "COELBA, 1.000 linhas (jan/2024)",
     "🔶 Não filtrado", "DscDetalheMercado: energia injetada/compensada (auxiliar)"),
    ("ANEEL_PeD", "ANEEL", "amostras/aneel_4.4_ped_projetos_COELBA.csv",
     "Projetos de P&D regulado: situação, tema, fase de inovação, produto, custos", "99 projetos COELBA, 2009–2026",
     "✅ Filtrado (dataset inteiro é CT&I)", "SigFasInovacaoProjeto; SigTemaProjeto; SigTipoProdutoProjeto"),
    ("ANEEL_DECFEC_Conjunto", "ANEEL", "amostras/aneel_4.5_decfec_conjunto_BA_2025.csv",
     "DEC e FEC anuais por conjunto elétrico, com limite regulatório", "211 conjuntos COELBA, 2025", "⬜ Sem variável", "—"),
    ("ANEEL_DECFEC_Municipio", "ANEEL", "amostras/aneel_4.5_decfec_municipio_BA_2025.csv",
     "DEC e FEC por município (média dos conjuntos, ponderada por consumidores)", "415 municípios, 2025", "⬜ Sem variável", "—"),
    ("ANEEL_INDQUAL", "ANEEL", "amostras/aneel_4.5_indqual_conjunto_municipio_BA.csv",
     "De-para conjunto elétrico ↔ município", "211 conjuntos ativos em 2025", "⬜ Sem variável", "—"),
    ("BNDES_Nao_Automaticas", "BNDES", "amostras/bndes_5.x_nao_automaticas_BA_amostra.csv",
     "Operações diretas e indiretas não automáticas, com colunas de classificação CT&I", "15 CT&I + 5 auxiliar + 10 fora",
     "🏷️ Classificado", "inovacao; instrumento_financeiro; fonte_de_recurso_desembolsos → cti_nivel"),
    ("BNDES_Indiretas_Automaticas", "BNDES", "amostras/bndes_5.x_indiretas_automaticas_BA_amostra.csv",
     "Operações indiretas automáticas (FINAME, BNDES Automático), com colunas CT&I", "15 CT&I + 5 auxiliar + 10 fora",
     "🏷️ Classificado", "inovacao; instrumento_financeiro; fonte_de_recurso_desembolsos → cti_nivel"),
    ("BNDES_CTI_Municipio_Ano", "BNDES", "amostras/bndes_5.5_cti_por_municipio_ano_BA.csv",
     "Operações CT&I somadas por município e ano", "Completo: 76 municípios × anos, 2002–2026", "✅ Filtrado",
     "cti_nivel = CT&I (regra explícita)"),
    ("SICONFI_Receita", "SICONFI", "amostras/siconfi_6.2_dca_anexoIC_receita_salvador_2024_amostra.csv",
     "Receitas por conta (DCA Anexo I-C)", "Salvador 2024, 40 linhas", "⬜ Sem variável", "—"),
    ("SICONFI_RREO_CTI", "SICONFI", "amostras/siconfi_6.3_rreo_anexo02_salvador_2025_cti.csv",
     "Despesa por função/subfunção (RREO Anexo 02)", "Salvador 2025, 6º bimestre", "✅ Filtrado (texto da conta)",
     "conta contém Ciência, Tecnol ou TOTAL"),
    ("SICONFI_DCA_CTI", "SICONFI", "amostras/siconfi_6.3_6.4_dca_IE_BA_2024_ciencia_tecnologia_por_municipio.csv",
     "Despesa liquidada por município: total, Função 19, subfunções 571–573 e 126", "415 municípios, 2024",
     "🏷️ Classificado em colunas", "Função 19; subfunções 571/572/573; 126 (auxiliar)"),
    ("ANATEL_Densidade", "ANATEL", "amostras/anatel_7.1_densidade_scm_BA_amostra.csv",
     "Acessos de banda larga fixa por 100 domicílios, por município e mês", "5 municípios × 2021, 2022, 2025, 2026",
     "⬜ Sem variável", "—"),
    ("ANATEL_Acessos", "ANATEL", "amostras/anatel_acessos_bl_fixa_2026_BA_amostra.csv",
     "Acessos por prestadora, tecnologia, meio de acesso e velocidade (jan–ago/2026)", "40 maiores linhas, municípios de referência",
     "🔶 Não filtrado", "Meio de Acesso = Fibra; Faixa > 34Mbps (auxiliar)"),
    ("ANATEL_Cobertura_Movel", "ANATEL", "amostras/anatel_7.2_cobertura_movel_4g_BA.csv",
     "Cobertura de telefonia móvel 4G: % de área coberta, % de moradores e % de domicílios cobertos",
     "417 municípios (completo BA) × Tecnologia 4G (Todas)", "⬜ Sem variável", "—"),
]
FILTROS = ("Filtros_CTI", "Todas", "categorias_cti/filtros_cti_consolidado.csv",
           "Lista única de filtros CT&I: fonte, campo, operador, valores, nível e justificativa", "—", "—", "—")

# Colunas que são códigos e devem continuar texto (zeros à esquerda, CNPJ)
TEXTO = re.compile(r"cnpj|cpf|^CO_CINE|^cod_conta$|^NumCNPJ|^IdeNucleo|^CodCEG", re.I)
NUM_PONTO = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")
NUM_VIRG = re.compile(r"^-?(\d+)?,\d+$")

AZUL, CINZA = "1F3864", "F2F2F2"
CORES_FILTRO = {"✅": "E2F0D9", "🏷": "DCE6F2", "🔶": "FFF2CC"}
F = "Arial"
borda = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def valor(col, v):
    if v is None or v == "":
        return v
    v = v.strip('"') if v.startswith('"') else v  # o INEP grava rótulos CINE entre aspas
    if TEXTO.search(col):
        return v
    if NUM_PONTO.match(v):
        n = float(v)
        return int(n) if n.is_integer() and "." not in v else n
    if NUM_VIRG.match(v):
        return float(v.replace(",", "."))
    if v in ("True", "False"):
        return v == "True"
    return v


def cabecalho(ws, linha, n):
    for c in range(1, n + 1):
        cel = ws.cell(linha, c)
        cel.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cel.fill = PatternFill("solid", fgColor=AZUL)
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        cel.border = borda


def larguras(ws, inicio=1, maximo=45):
    for col in ws.iter_cols(min_row=inicio):
        tam = max((len(str(c.value)) for c in col[:300] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(tam + 2, 9), maximo)


def aba_dados(wb, aba, arquivo, descricao):
    ws = wb.create_sheet(aba)
    with open(ROOT / arquivo, encoding="utf-8-sig") as f:
        linhas = list(csv.reader(f))
    cols = linhas[0]
    ws.append(["← LEIA-ME", descricao, f"Origem: {arquivo}"])
    ws["A1"].hyperlink = "#'LEIA-ME'!A1"
    ws["A1"].font = Font(name=F, color="0563C1", underline="single", size=9)
    for c in ("B1", "C1"):
        ws[c].font = Font(name=F, italic=True, color="595959", size=9)
    ws.append(cols)
    cabecalho(ws, 2, len(cols))
    for l in linhas[1:]:
        ws.append([valor(c, v) for c, v in zip(cols, l)])
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.font = Font(name=F, size=9)
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"
            elif isinstance(c.value, int) and not re.search(r"^(ano|Ano|NU_ANO|codigo|cod_|CO_|Cod|Ide|municipio_codigo|Código)", ws.cell(2, c.column).value or ""):
                c.number_format = "#,##0"
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{ws.max_row}"
    larguras(ws, inicio=2)
    return len(linhas) - 1, len(cols)


wb = Workbook()
leia = wb.active
leia.title = "LEIA-ME"
info = {}
for a in ABAS + [FILTROS]:
    info[a[0]] = aba_dados(wb, a[0], a[2], a[3])

# --- LEIA-ME
leia["A1"] = "TERRITÓRIOS SECTI — Amostras de dados e filtros CT&I (Bahia)"
leia["A1"].font = Font(name=F, bold=True, size=14, color=AZUL)
textos = [
    "Cada aba traz a amostra de um dataset citado no resumo executivo. A linha 1 de cada aba repete a descrição e o arquivo de origem; a linha 2 é o cabeçalho, com filtro.",
    "A coluna \"Filtro CT&I\" diz se as linhas da aba já foram selecionadas por CT&I (✅), se trazem uma coluna que classifica cada linha (🏷️), se não foram filtradas mas têm variável CT&I (🔶) ou se não há variável CT&I (⬜).",
    "A aba Filtros_CTI é a regra oficial de recorte CT&I do painel. Clique no nome de uma aba para abri-la.",
]
for i, t in enumerate(textos, start=2):
    leia.cell(i, 1, t).font = Font(name=F, size=10)
    leia.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
    leia.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
    leia.row_dimensions[i].height = 30
cab = ["Aba", "Fonte", "O que retorna", "Recorte da amostra", "Filtro CT&I", "Variáveis CT&I", "Linhas", "Colunas", "Arquivo de origem"]
L0 = 6
for j, c in enumerate(cab, 1):
    leia.cell(L0, j, c)
cabecalho(leia, L0, len(cab))
for k, (aba, fonte, arq, desc, rec, filt, var) in enumerate([FILTROS] + ABAS, start=L0 + 1):
    q = f"'{aba}'"
    vals = [aba, fonte, desc, rec, filt, var,
            f"=COUNTA({q}!A:A)-2",                      # linhas de dados (descontando descrição e cabeçalho)
            f"=COUNTA({q}!2:2)", arq]
    for j, v in enumerate(vals, 1):
        cel = leia.cell(k, j, v)
        cel.font = Font(name=F, size=9)
        cel.border = borda
        cel.alignment = Alignment(wrap_text=True, vertical="top")
        cor = next((h for s, h in CORES_FILTRO.items() if str(filt).startswith(s)), None)
        cel.fill = PatternFill("solid", fgColor=cor or (CINZA if k % 2 else "FFFFFF"))
    leia.cell(k, 1).hyperlink = f"#{q}!A1"
    leia.cell(k, 1).font = Font(name=F, size=9, color="0563C1", underline="single")
    leia.cell(k, 7).number_format = "#,##0"
for col, w in zip("ABCDEFGHI", [26, 9, 48, 34, 22, 36, 9, 9, 46]):
    leia.column_dimensions[col].width = w
leia.freeze_panes = f"A{L0 + 1}"
nota = L0 + len(ABAS) + 3
leia.cell(nota, 1, "Fontes e datas: coleta de 01 a 07/10/2026; detalhes, regras e ressalvas em RESULTADOS_AMOSTRAGEM.md. "
          "Linhas e Colunas são fórmulas que contam o conteúdo de cada aba.").font = Font(name=F, italic=True, size=9, color="595959")

wb.save(SAIDA)
print(SAIDA, {k: v[0] for k, v in info.items()})
